"""Gera os .docx mínimos de teste do perfil de ingestão.

    python tests/make_fixtures.py

Saída em `tests/fixtures/` (arquivos pequenos, versionados):
- `mini-muriae.docx`: formato do preset (separador `***`, alternativas `A)`).
- `mini-alt.docx`: formato alternativo (separador `---`, alternativas `1)`),
  que o preset NÃO segmenta — só o `mini-alt.json` segmenta.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import RGBColor

FIXTURES = Path(__file__).resolve().parent / "fixtures"
RED = RGBColor(0xFF, 0x00, 0x00)


def _red(paragraph) -> None:
    for run in paragraph.runs:
        run.font.color.rgb = RED


def make_muriae(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Descritor 1: Numeros e operacoes")
    doc.add_paragraph(
        "Com este descritor avalia-se a habilidade de o aluno resolver problemas."
    )
    doc.add_paragraph("ATIVIDADES DO DESCRITOR 1")
    doc.add_paragraph("(Prova Brasil). Quanto e 2 + 2?")
    for letter, text in [("A", "3"), ("B", "4"), ("C", "5"), ("D", "6")]:
        p = doc.add_paragraph(f"{letter}) {text}")
        if letter == "B":
            _red(p)
    doc.add_paragraph("********************")
    doc.add_paragraph("Quanto e 3 + 3?")
    for letter, text in [("A", "5"), ("B", "7"), ("C", "6"), ("D", "9")]:
        p = doc.add_paragraph(f"{letter}) {text}")
        if letter == "C":
            _red(p)
    doc.save(path)


def make_alt(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Descritor 1: Numeros e operacoes")
    doc.add_paragraph(
        "Com este descritor avalia-se a habilidade de o aluno resolver problemas."
    )
    doc.add_paragraph("ATIVIDADES DO DESCRITOR 1")
    doc.add_paragraph("(Prova Brasil). Quanto e 2 + 2?")
    for number, text in [("1", "3"), ("2", "4"), ("3", "5"), ("4", "6")]:
        p = doc.add_paragraph(f"{number}) {text}")
        if number == "2":
            _red(p)
    doc.add_paragraph("---")
    doc.add_paragraph("Quanto e 3 + 3?")
    for number, text in [("1", "5"), ("2", "7"), ("3", "6"), ("4", "9")]:
        p = doc.add_paragraph(f"{number}) {text}")
        if number == "3":
            _red(p)
    doc.save(path)


if __name__ == "__main__":
    FIXTURES.mkdir(parents=True, exist_ok=True)
    make_muriae(FIXTURES / "mini-muriae.docx")
    make_alt(FIXTURES / "mini-alt.docx")
    print(f"fixtures em {FIXTURES}")
